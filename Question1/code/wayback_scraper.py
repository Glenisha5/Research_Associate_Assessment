#!/usr/bin/env python3
"""
wayback_scraper.py
==================
Collect, for every incubator in an input list:

  (i)  the HOMEPAGE and (ii) the ABOUT-US page of the CURRENT live website, and
  (iii) the same two pages as ARCHIVED on the Internet Archive's Wayback Machine,
        using the EARLIEST available capture in EACH calendar month,
        walking backwards from --end (default 2026-09) to --start (default 2022-01).

Everything is saved as raw HTML plus a `manifest.csv` that records, for every
(incubator, page, month), what happened and where the file is.

Quick start
-----------
    pip install -r requirements.txt
    python wayback_scraper.py --input incubators.csv --out output --contact you@university.edu

See README.md for the full documentation.
"""
from __future__ import annotations

import argparse
import calendar
import csv
import datetime as dt
import hashlib
import logging
import re
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Optional
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

# ----------------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------------
# Wayback "CDX" index: lets us ask "which captures of URL X exist between date A and B?"
CDX_API = "https://web.archive.org/cdx/search/cdx"

# Wayback replay URL. The `id_` suffix after the timestamp asks for the ORIGINAL bytes
# exactly as archived (no Wayback toolbar, no rewritten links) - best for research.
WAYBACK_RAW = "https://web.archive.org/web/{ts}id_/{url}"

# Paths tried (in this order) when no about-page URL is supplied or discoverable.
ABOUT_PATHS = [
    "about-us", "about-us/", "about", "about/", "about.html", "aboutus",
    "about-us.html", "about-us.php", "about.php", "who-we-are", "who-we-are/", "overview",
]

# Used to spot an "about" link inside a homepage (matches link text OR URL path).
ABOUT_LINK_RE = re.compile(r"\babout(\s|-|_)?(us)?\b|who we are|our story", re.I)

# Strict: only captures where the server answered 200 OK.
# Relaxed: also accept redirects (301/302/303/307/308) and Wayback "revisit" records ("-").
STATUS_STRICT = "statuscode:200"
STATUS_RELAXED = r"statuscode:(200|30[12378]|-)"

MANIFEST_FIELDS = [
    "incubator", "page_type", "month", "status", "snapshot_timestamp",
    "requested_url", "archived_url", "file", "bytes", "sha256", "note",
]


# ----------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------
def month_labels(start: str, end: str) -> list[str]:
    """Return ['YYYY-MM', ...] from `end` back to `start` (newest first), inclusive."""
    sy, sm = map(int, start.split("-"))
    ey, em = map(int, end.split("-"))
    labels, (y, m) = [], (ey, em)
    while (y, m) >= (sy, sm):
        labels.append(f"{y}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return labels


def slugify(text: str) -> str:
    """'NSRCEL (IIM Bangalore)' -> 'nsrcel_iim_bangalore' (safe folder name)."""
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def save(path: Path, content: bytes) -> tuple[int, str]:
    """Write bytes to disk (creating folders); return (size, sha256) for the manifest."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    return len(content), hashlib.sha256(content).hexdigest()


def load_rows(path: Path) -> list[dict]:
    """Read the incubator list from .csv OR .xlsx.
    Required columns: name, homepage_url.  Optional: about_url."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook  # imported lazily: only needed for Excel input
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip().lower() if h else "" for h in next(it)]
        rows = [dict(zip(header, ["" if v is None else str(v).strip() for v in r])) for r in it]
    else:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            rows = [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
                    for r in csv.DictReader(fh)]
    rows = [r for r in rows if r.get("name") and r.get("homepage_url")]
    if not rows:
        sys.exit(f"No usable rows in {path}. Need columns: name, homepage_url[, about_url].")
    return rows


# ----------------------------------------------------------------------------
# HTTP layer: User-Agent, polite delay, retry with exponential back-off
# ----------------------------------------------------------------------------
class Http:
    def __init__(self, contact: str, delay: float = 1.0, retries: int = 5, timeout: int = 60):
        self.session = requests.Session()
        # Identify yourself: archive.org is friendlier to bots that say who they are.
        self.session.headers["User-Agent"] = f"IncubatorResearchBot/2.0 (academic research; contact: {contact})"
        self.delay, self.retries, self.timeout = delay, retries, timeout

    def get(self, url: str, **kw) -> Optional[requests.Response]:
        """GET with retries. Returns the Response on HTTP 200, otherwise None.
        429/5xx and network errors are retried (wait 4,8,16,... capped at 120 s);
        other codes (404, 403...) are final and not retried."""
        for attempt in range(1, self.retries + 1):
            try:
                time.sleep(self.delay)  # politeness delay before EVERY request
                r = self.session.get(url, timeout=self.timeout, **kw)
                if r.status_code == 200:
                    return r
                if r.status_code in (429, 500, 502, 503, 504):
                    wait = min(2 ** attempt * 2, 120)
                    logging.warning("HTTP %s for %s - retry %d/%d in %ds",
                                    r.status_code, url, attempt, self.retries, wait)
                    time.sleep(wait)
                    continue
                logging.info("HTTP %s for %s (not retrying)", r.status_code, url)
                return None
            except requests.RequestException as e:
                wait = min(2 ** attempt * 2, 120)
                logging.warning("Network error %s for %s - retry %d/%d in %ds",
                                e, url, attempt, self.retries, wait)
                time.sleep(wait)
        return None


# ----------------------------------------------------------------------------
# Wayback logic
# ----------------------------------------------------------------------------
def cdx_monthly(http: Http, url: str, start: str, end: str, relaxed: bool) -> Optional[dict[str, tuple[str, str]]]:
    """ONE index query returns the earliest capture of `url` in every month of the range.

    How: the CDX API lists captures in ascending time order; `collapse=timestamp:6`
    keeps only the first row per YYYYMM (6 leading timestamp digits) - i.e. the
    earliest capture of each month. This replaces ~57 per-month queries with 1.

    Returns {'2024-03': (timestamp, original_url), ...}, or None if the query failed
    (so callers can tell "no captures" ({}) apart from "couldn't ask" (None)).
    """
    sy, sm = map(int, start.split("-"))
    ey, em = map(int, end.split("-"))
    params = {
        "url": url,
        "matchType": "exact",
        "from": f"{sy}{sm:02d}01000000",
        "to": f"{ey}{em:02d}{calendar.monthrange(ey, em)[1]:02d}235959",  # true last day of month
        "output": "json",
        "fl": "timestamp,original,statuscode",
        "filter": STATUS_RELAXED if relaxed else STATUS_STRICT,
        "collapse": "timestamp:6",
    }
    r = http.get(CDX_API, params=params)
    if r is None:
        return None
    try:
        rows = r.json()
    except ValueError:
        return None
    caps: dict[str, tuple[str, str]] = {}
    for ts, original, _status in rows[1:]:  # rows[0] is the header row
        caps.setdefault(f"{ts[:4]}-{ts[4:6]}", (ts, original))  # keep the first (earliest)
    return caps


def fetch_snapshot(http: Http, ts: str, original: str) -> Optional[bytes]:
    """Download the raw archived HTML for one capture."""
    r = http.get(WAYBACK_RAW.format(ts=ts, url=original))
    return r.content if r is not None else None


def discover_about_url(html: bytes, base_url: str) -> Optional[str]:
    """Look through a homepage's <a> tags for an 'About' link on the same site.
    If several match, prefer the shortest (most generic) path."""
    soup = BeautifulSoup(html, "html.parser")
    host = urlparse(base_url).netloc.replace("www.", "")
    best: Optional[str] = None
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        full = urljoin(base_url, href)
        if urlparse(full).netloc.replace("www.", "") != host:
            continue  # ignore external links
        if ABOUT_LINK_RE.search(a.get_text(" ", strip=True)) or ABOUT_LINK_RE.search(urlparse(full).path):
            if best is None or len(urlparse(full).path) < len(urlparse(best).path):
                best = full
    return best


def fallback_about_urls(homepage: str) -> list[str]:
    """Guess likely about-page URLs. Tries them relative to the site root AND to the
    homepage path (some 'homepages' are sub-pages such as /research-park/incubation-cell)."""
    p = urlparse(homepage)
    root = f"{p.scheme}://{p.netloc}/"
    hp = homepage if homepage.endswith("/") else homepage + "/"
    bases = [root] if hp == root else [root, hp]
    out: list[str] = []
    for base in bases:
        for path in ABOUT_PATHS:
            u = urljoin(base, path)
            if u not in out:
                out.append(u)
    return out


# ----------------------------------------------------------------------------
# Manifest writer (also keeps tallies for the end-of-run summary)
# ----------------------------------------------------------------------------
class Manifest:
    def __init__(self, path: Path):
        self.fh = open(path, "w", newline="", encoding="utf-8")
        self.w = csv.DictWriter(self.fh, fieldnames=MANIFEST_FIELDS)
        self.w.writeheader()
        self.counts: Counter = Counter()

    def row(self, **kw) -> None:
        self.counts[(kw["page_type"], kw["status"])] += 1
        self.w.writerow({f: kw.get(f, "") for f in MANIFEST_FIELDS})
        self.fh.flush()  # flush every row so a crash never loses the log

    def close(self) -> None:
        self.fh.close()


# ----------------------------------------------------------------------------
# Main per-incubator workflow
# ----------------------------------------------------------------------------
def process_incubator(row: dict, http: Http, args, out: Path, labels: list[str], mf: Manifest) -> None:
    name = row["name"]
    slug = slugify(name)
    homepage = row["homepage_url"]
    about_known = row.get("about_url") or None  # optional manual override from the spreadsheet
    logging.info("=== %s (%s)", name, homepage)

    def log(page, month, status, **kw):
        mf.row(incubator=name, page_type=page, month=month, status=status, **kw)

    def existing(page: str, label: str) -> Optional[Path]:
        """Resume support: has this month's file already been saved in an earlier run?"""
        hits = sorted((out / slug / page).glob(f"{label}_*.html"))
        return hits[0] if hits else None

    def log_existing(page: str, label: str, f: Path) -> bytes:
        data = f.read_bytes()
        log(page, label, "skipped_exists", file=str(f), bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
        return data

    def download_first_working(page: str, label: str, options: list[tuple[str, str, str]]) -> Optional[bytes]:
        """options = [(requested_url, timestamp, original_url), ...] in priority order.
        Save the first one that downloads successfully."""
        for req, ts, original in options:
            content = fetch_snapshot(http, ts, original)
            if content:
                f = out / slug / page / f"{label}_{ts}.html"
                size, sha = save(f, content)
                log(page, label, "ok", snapshot_timestamp=ts, requested_url=req,
                    archived_url=WAYBACK_RAW.format(ts=ts, url=original), file=str(f), bytes=size, sha256=sha)
                return content
        log(page, label, "error", note="capture(s) listed in index but download failed")
        return None

    # ---------------- (a) CURRENT live pages ----------------------------------
    live_about: Optional[str] = None
    if not args.skip_current:
        today = dt.date.today().strftime("%Y%m%d")

        def fetch_live(page: str, url: str) -> Optional[bytes]:
            f = out / slug / page / f"current_{today}.html"
            if f.exists():  # already fetched today -> resume
                return log_existing(page, "current", f)
            r = http.get(url)
            if r is None:
                log(page, "current", "error", requested_url=url, note="live fetch failed")
                return None
            size, sha = save(f, r.content)
            log(page, "current", "ok", requested_url=url, file=str(f), bytes=size, sha256=sha)
            return r.content

        live_home = fetch_live("homepage", homepage)
        live_about = about_known or (discover_about_url(live_home, homepage) if live_home else None)
        if live_about:
            fetch_live("about", live_about)
        else:
            log("about", "current", "missing", note="no about link found on live homepage")

    # ---------------- (b) ARCHIVED homepage: earliest capture per month ---------
    home_html: dict[str, bytes] = {}  # kept in memory: used later to discover about links
    todo = []
    for label in labels:
        f = existing("homepage", label)
        if f:
            home_html[label] = log_existing("homepage", label, f)
        else:
            todo.append(label)
    if todo:
        caps = cdx_monthly(http, homepage, args.start, args.end, args.relaxed_status)
        for label in todo:
            if caps is None:
                log("homepage", label, "error", requested_url=homepage, note="CDX index query failed")
            elif label in caps:
                ts, original = caps[label]
                data = download_first_working("homepage", label, [(homepage, ts, original)])
                if data:
                    home_html[label] = data
            else:
                log("homepage", label, "no_snapshot", requested_url=homepage,
                    note="no capture this month (with the chosen status filter)")

    # ---------------- (c) ARCHIVED about page: earliest capture per month -------
    todo = []
    for label in labels:
        f = existing("about", label)
        if f:
            log_existing("about", label, f)
        else:
            todo.append(label)

    if todo:
        options: dict[str, list[tuple[str, str, str]]] = {l: [] for l in todo}
        remaining = set(todo)
        tried: dict[str, Optional[dict]] = {}

        # Candidate order = manual URL, live-discovered URL, then guessed paths.
        # Candidates are queried lazily: stop as soon as every month is covered.
        candidates: list[str] = []
        for c in [about_known, live_about] + fallback_about_urls(homepage):
            if c and c not in candidates:
                candidates.append(c)
        for cand in candidates:
            if not remaining:
                break
            caps = cdx_monthly(http, cand, args.start, args.end, args.relaxed_status)
            tried[cand] = caps
            if not caps:
                continue
            for label in list(remaining):
                if label in caps:
                    options[label].append((cand, *caps[label]))
                    remaining.discard(label)

        # Last resort for still-uncovered months: read that month's ARCHIVED homepage
        # and follow whatever "About" link it contained at the time (site may have been redesigned).
        for label in sorted(remaining, reverse=True):
            html = home_html.get(label)
            d = discover_about_url(html, homepage) if html else None
            if not d:
                continue
            if d not in tried:
                tried[d] = cdx_monthly(http, d, args.start, args.end, args.relaxed_status)
            if tried[d] and label in tried[d]:
                options[label].append((d, *tried[d][label]))

        for label in todo:
            if options[label]:
                download_first_working("about", label, options[label])
            else:
                log("about", label, "no_snapshot", note=f"checked {len(tried)} candidate URLs; none captured this month")


# ----------------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------------
def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default="incubators.csv", help="CSV or XLSX with columns name, homepage_url[, about_url]")
    ap.add_argument("--out", default="output", help="output directory (default: output)")
    ap.add_argument("--contact", required=True, help="your email; sent in the User-Agent header")
    ap.add_argument("--start", default="2022-01", help="oldest month, YYYY-MM (default 2022-01)")
    ap.add_argument("--end", default="2026-09", help="newest month, YYYY-MM (default 2026-09)")
    ap.add_argument("--delay", type=float, default=1.0, help="seconds between HTTP requests (default 1.0)")
    ap.add_argument("--only", help="only run incubators whose name contains this text (testing)")
    ap.add_argument("--skip-current", action="store_true", help="do not fetch the live (current) pages")
    ap.add_argument("--relaxed-status", action="store_true",
                    help="also accept redirect (30x) and revisit captures, not just HTTP 200")
    a = ap.parse_args(argv)

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.FileHandler(out / "run.log", encoding="utf-8"),
                                  logging.StreamHandler(sys.stdout)])

    rows = load_rows(Path(a.input))
    if a.only:
        rows = [r for r in rows if a.only.lower() in r["name"].lower()]
    labels = month_labels(a.start, a.end)
    logging.info("%d incubators x %d months (%s back to %s)", len(rows), len(labels), a.end, a.start)

    http = Http(a.contact, delay=a.delay)
    mf = Manifest(out / "manifest.csv")
    try:
        for row in rows:
            process_incubator(row, http, a, out, labels, mf)
    finally:
        mf.close()

    logging.info("Done. Manifest: %s", out / "manifest.csv")
    for (page, status), n in sorted(mf.counts.items()):
        logging.info("  %-9s %-15s %d", page, status, n)


if __name__ == "__main__":
    main()
