"""
wipo_merge.py  -  merge raw WIPO PATENTSCOPE exports with the startup sheet.

Usage:
    python wipo_merge.py <raw_wipo_dir> <ranked_startup_xlsx> <search_log_xlsx> <output_xlsx>

Steps (see Matching_Process.md for the written description):
 1. Read every .xls/.xlsx/.csv export in raw_wipo/. PATENTSCOPE puts 5 header lines
    (Time, Query, SortedBy) above the table, so the real header row is detected.
 2. Split the Applicants cell on ';' into separate parties (joint filings, inventors).
 3. Normalise each party name: lower-case, strip punctuation, drop legal suffixes
    (pvt, private, ltd, limited, pte, b.v., inc ...).
 4. Match against the alias table (ALIASES) built from the startups' registered names:
      exact normalised match -> 'exact-normalised'
      similarity >= 0.90      -> 'fuzzy-auto'
      0.80 - 0.90             -> 'review'
      below 0.80              -> rejected
 5. Cross-check that the matched startup equals the startup the file was searched for.
 6. Write: Patents_Merged, Rejected_Rows, Startup_Summary (joined to the startup sheet,
    with live formulas), Alias_Table, File_Checks.
"""
import re, sys, glob, os, difflib
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# startup name (as in the ranked sheet) -> registered-name aliases seen/used in WIPO.
ALIASES = {
    "Ola Electric":             ["Ola Electric Mobility"],
    "Ather Energy":             ["Ather Energy"],
    "River":                    ["River Mobility"],
    "SUN Mobility":             ["SUN Mobility"],
    "Ultraviolette Automotive": ["Ultraviolette Automotive"],
    "Matter":                   ["Matter Motor Works"],
    "Ecozen Solutions":         ["Ecozen Solutions"],
    "Smartron":                 ["Smartron India"],
    "Carbon Clean Solutions":   ["Carbon Clean Solutions"],
    "Chakr Innovation":         ["Chakr Innovation"],
}
# file-number prefix -> startup the file was searched for (from the file name)
FILE_STARTUP = {"01": "Ola Electric", "03": "Ather Energy", "07": "River", "10": "SUN Mobility",
                "11": "Ultraviolette Automotive", "15": "Matter", "16": "Ecozen Solutions",
                "17": "Smartron", "18": "Carbon Clean Solutions", "19": "Chakr Innovation"}
SUFFIX = r"\b(pvt|private|ltd|limited|pte|bv|inc|llc|llp|corp|corporation|gmbh|co)\b"

def norm(s):
    s = re.sub(r"\s+", " ", str(s).lower().replace(".", " "))   # 'PVT.LTD.' -> 'pvt ltd'
    s = re.sub(r"\bb v\b", "bv", s)
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    s = re.sub(SUFFIX, " ", s)
    return re.sub(r"\s+", " ", s).strip()

def read_export(path):
    raw = (pd.read_csv(path, header=None, dtype=str, on_bad_lines="skip") if path.endswith(".csv")
           else pd.read_excel(path, header=None, dtype=str))
    hdr = next(i for i in range(min(15, len(raw)))
               if "Application Number" in [str(v).strip() for v in raw.iloc[i]])
    meta = {"time": str(raw.iloc[1, 1]), "query": str(raw.iloc[2, 1])}
    df = raw.iloc[hdr + 1:].copy()
    df.columns = [str(v).strip() for v in raw.iloc[hdr]]
    return df.dropna(how="all").rename(columns={"I P C": "IPC"}), meta

def match_party(party, alias_norm):
    n = norm(party)
    if not n: return None, "", 0.0
    if n in alias_norm: return alias_norm[n], "exact-normalised", 1.0
    best = max(alias_norm, key=lambda k: difflib.SequenceMatcher(None, n, k).ratio())
    sc = round(difflib.SequenceMatcher(None, n, best).ratio(), 2)
    if sc >= 0.90: return alias_norm[best], "fuzzy-auto", sc
    if sc >= 0.80: return alias_norm[best], "review", sc
    return None, "no-match", sc

def main(raw_dir, ranked, log, out):
    for label, path in (("raw_wipo folder", raw_dir), ("startup sheet", ranked), ("search log", log)):
        if not os.path.exists(path):
            sys.exit(f"ERROR: {label} not found: '{path}'\n"
                     f"Current folder is {os.getcwd()}. Put the files in this folder (or give full paths) and run again.")
    if not glob.glob(os.path.join(raw_dir, "*.xls*")) and not glob.glob(os.path.join(raw_dir, "*.csv")):
        sys.exit(f"ERROR: no .xls/.xlsx/.csv files inside '{raw_dir}'.")
    alias_norm = {norm(a): s for s, al in ALIASES.items() for a in al}
    matched, rejected, checks = [], [], []
    for f in sorted(glob.glob(os.path.join(raw_dir, "*"))):
        if not f.lower().endswith((".xls", ".xlsx", ".csv")): continue
        df, meta = read_export(f); base = os.path.basename(f)
        searched = FILE_STARTUP.get(base[:2])
        n_m = n_r = 0
        for _, r in df.iterrows():
            parties = [p.strip() for p in str(r["Applicants"]).split(";") if p.strip()]
            res = [(p,) + match_party(p, alias_norm) for p in parties]
            hits = [x for x in res if x[1]]
            base_row = {"Source file": base, "Query": meta["query"], "Export time": meta["time"],
                        "Application Id": r["Application Id"], "Application Number": r["Application Number"],
                        "Application Date": r["Application Date"], "Publication Number": r["Publication Number"],
                        "Publication Date": r["Publication Date"], "Country": r["Country"],
                        "Title": r["Title"], "IPC": r["IPC"], "Applicants (as in WIPO)": r["Applicants"]}
            if hits:
                p, s, how, sc = max(hits, key=lambda x: x[3])
                others = "; ".join(x[0] for x in res if x[0] != p)
                flag = []
                if how == "review": flag.append("fuzzy match, check by hand")
                if s != searched: flag.append(f"startup differs from file ({searched})")
                if others: flag.append("other parties listed")
                matched.append({**base_row, "Startup": s, "Matched party": p, "Match method": how,
                                "Match score": sc, "Other parties on filing": others,
                                "Review flag": "; ".join(flag)}); n_m += 1
            else:
                best = max(res, key=lambda x: x[3]) if res else ("", None, "no-match", 0)
                rejected.append({**base_row, "Searched for": searched, "Best similarity": best[3],
                                 "Reason": "applicant name does not match the startup's registered name"}); n_r += 1
        checks.append({"File": base, "Searched for": searched, "Query in file": meta["query"],
                       "Rows in file": len(df), "Matched": n_m, "Rejected": n_r,
                       "Duplicate publication numbers": int(df["Publication Number"].duplicated().sum())})
    M = pd.DataFrame(matched); R = pd.DataFrame(rejected); C = pd.DataFrame(checks)
    M["Publication year"] = pd.to_datetime(M["Publication Date"], format="%d.%m.%Y", errors="coerce").dt.year
    rank = openpyxl.load_workbook(ranked, data_only=True)["Startups"]
    lg = {r[1]: r for r in openpyxl.load_workbook(log, data_only=True).active.iter_rows(min_row=2, values_only=True) if r[1]}
    top = [r for r in rank.iter_rows(min_row=2, values_only=True) if r[1] and r[18] == "Yes"]
    top.sort(key=lambda r: r[15])
    write_book(out, M, R, C, top, lg)
    print(f"{len(M)} matched rows, {len(R)} rejected rows -> {out}")

# ---------------------------------------------------------------- workbook
F = "Arial"
HF = Font(name=F, bold=True, color="FFFFFF", size=10); HFILL = PatternFill("solid", fgColor="1F4E3D")
BODY = Font(name=F, size=10); BLUE = Font(name=F, size=10, color="0000FF"); BOLD = Font(name=F, size=10, bold=True)
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def put_df(ws, df, widths):
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(1, j, c); x.font, x.fill, x.border = HF, HFILL, BOX
        x.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            v = None if pd.isna(v) else v
            x = ws.cell(i, j, v); x.font = BODY
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

def write_book(out, M, R, C, top, lg):
    wb = openpyxl.Workbook()
    # --- Patents_Merged
    ws = wb.active; ws.title = "Patents_Merged"
    cols = ["Startup", "Source file", "Query", "Application Id", "Application Number", "Application Date",
            "Publication Number", "Publication Date", "Publication year", "Country", "Title", "IPC",
            "Applicants (as in WIPO)", "Matched party", "Match method", "Match score",
            "Other parties on filing", "Review flag", "Export time"]
    put_df(ws, M[cols], [22, 28, 28, 16, 20, 12, 20, 12, 8, 8, 60, 30, 44, 36, 16, 8, 40, 36, 18])
    nM = len(M) + 1
    # --- Rejected
    ws = wb.create_sheet("Rejected_Rows")
    rc = ["Searched for", "Source file", "Query", "Application Id", "Publication Number", "Publication Date",
          "Country", "Title", "Applicants (as in WIPO)", "Best similarity", "Reason"]
    if len(R): put_df(ws, R[rc], [22, 28, 28, 16, 20, 12, 8, 60, 44, 10, 50])
    else: ws["A1"] = "No rows rejected."
    # --- Startup_Summary
    ws = wb.create_sheet("Startup_Summary", 0)
    h = ["Final Rank", "Startup", "Category", "Founded", "HQ", "Total funding (USD M)", "Employees",
         "WIPO query used", "Result count on WIPO page", "Rows downloaded", "Matched patent records",
         "Rows rejected (unrelated applicant)", "Unique publication numbers", "First publication year",
         "Last publication year", "Records: IN office", "Records: WO (PCT)", "Records: other offices", "Search outcome / notes"]
    for j, c in enumerate(h, 1):
        x = ws.cell(1, j, c); x.font, x.fill, x.border = HF, HFILL, BOX
        x.alignment = Alignment(wrap_text=True, vertical="center")
    rngS = f"Patents_Merged!$A$2:$A${nM}"; rngP = f"Patents_Merged!$G$2:$G${nM}"
    rngY = f"Patents_Merged!$I$2:$I${nM}"; rngC = f"Patents_Merged!$J$2:$J${nM}"
    dl = {}
    for _, r in C.iterrows():
        dl[r["Searched for"]] = (r["Rows in file"], r["Query in file"])
    for i, t in enumerate(top, 2):
        name = t[1]; l = lg.get(name)
        got = dl.get(name)
        q = got[1] if got else (l[2] if l else "")
        page = l[4] if l else None
        note = ("Records found and matched." if got else
                (l[7] if l and l[7] else "No relevant WIPO results; no file downloaded."))
        vals = [t[15], name, t[10], t[5], t[4], t[12], t[7], q, page, got[0] if got else 0]
        for j, v in enumerate(vals, 1):
            x = ws.cell(i, j, v); x.font = BLUE if j in (6, 7, 9, 10) else BODY; x.border = BOX
        f = {11: f"=COUNTIF({rngS},B{i})", 12: f"=J{i}-K{i}",
             13: f"=IFERROR(SUMPRODUCT(({rngS}=B{i})/COUNTIFS({rngS},{rngS},{rngP},{rngP})),0)",
             14: f'=IF(K{i}=0,"",9999-SUMPRODUCT(MAX(({rngS}=B{i})*(9999-{rngY}))))', 15: f'=IF(K{i}=0,"",SUMPRODUCT(MAX(({rngS}=B{i})*{rngY})))',
             16: f'=COUNTIFS({rngS},B{i},{rngC},"IN")', 17: f'=COUNTIFS({rngS},B{i},{rngC},"WO")',
             18: f"=K{i}-P{i}-Q{i}"}
        for j, v in f.items():
            x = ws.cell(i, j, v); x.font = BODY; x.border = BOX
        x = ws.cell(i, 19, note); x.font = BODY; x.border = BOX; x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 6).number_format = "#,##0.0"; ws.cell(i, 13).number_format = "0"
    n = len(top) + 2
    ws.cell(n, 2, "Total").font = BOLD
    for j in (10, 11, 12, 16, 17, 18):
        c = get_column_letter(j); x = ws.cell(n, j, f"=SUM({c}2:{c}{n-1})"); x.font = BOLD
    ws.cell(n + 2, 2, "Blue = input from the startup sheet or the WIPO search log. Black = formula over Patents_Merged. 'Records' are WIPO patent records, "
                      "not unique inventions: one invention filed in several countries appears several times. The export carries no grant/kind code, "
                      "so filed vs granted cannot be separated.").font = Font(name=F, size=9, italic=True)
    for j, w in enumerate([8, 24, 20, 8, 22, 12, 10, 30, 12, 11, 12, 14, 12, 10, 10, 10, 10, 10, 60], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"; ws.row_dimensions[1].height = 54
    # --- Alias table
    ws = wb.create_sheet("Alias_Table")
    a = pd.DataFrame([(s, al, norm(al)) for s, al_ in ALIASES.items() for al in al_],
                     columns=["Startup", "Registered-name alias searched/matched", "Normalised form"])
    put_df(ws, a, [26, 40, 30])
    # --- File checks
    ws = wb.create_sheet("File_Checks")
    put_df(ws, C, [34, 24, 34, 12, 10, 10, 18])
    ws.cell(len(C) + 3, 1, "Rows in file = Matched + Rejected for every file. Compare 'Rows in file' with the count in WIPO_Search_Log.").font = Font(name=F, size=9, italic=True)
    wb.save(out)

if __name__ == "__main__":
    main(*sys.argv[1:5])